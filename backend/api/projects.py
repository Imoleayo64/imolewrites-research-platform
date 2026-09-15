from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from sqlalchemy.orm import Session
from typing import List, Optional
from pydantic import BaseModel
import io
import re
import docx
from fpdf import FPDF

from backend.crud.projects import create_project, get_projects, get_project, update_project, delete_project
from backend.database.session import get_db
from backend.models.project import ProjectCreate, ProjectResponse
from backend.services.auth_service import get_current_user
from backend.services.analytics_service import log_event
from backend.services.html_export_service import html_to_docx, html_to_markdown, html_to_plain_text
from backend.database.models import ChapterModel

router = APIRouter(
    prefix="/projects",
    tags=["Projects"]
)

class CitationExport(BaseModel):
    bibtex: Optional[str] = None
    formatted: Optional[dict] = None
    title: Optional[str] = None

class ExportRequest(BaseModel):
    title: str
    content: str
    format: str
    citations: List[CitationExport] = []

class ProjectUpdate(BaseModel):
    title: str
    content: str

class ProjectStatusUpdate(BaseModel):
    status: str  # draft, submitted, in_review, revisions, published
    target_journal: str | None = None

@router.put("/{project_id}", response_model=ProjectResponse)
def update_project_data(
    project_id: int, 
    update_data: ProjectUpdate, 
    db: Session = Depends(get_db), 
    current_user = Depends(get_current_user)
):
    updated_project = update_project(db, project_id=project_id, user_id=current_user.id, title=update_data.title, content=update_data.content)
    if not updated_project:
        raise HTTPException(status_code=404, detail="Project not found")
    word_count = len(update_data.content.split()) if update_data.content else 0
    log_event(db, user_id=current_user.id, event_type="project_save", value=word_count)
    return updated_project

@router.patch("/{project_id}/status", response_model=ProjectResponse)
def update_project_status(
    project_id: int,
    payload: ProjectStatusUpdate,
    db: Session = Depends(get_db),
    current_user = Depends(get_current_user)
):
    valid_statuses = {"draft", "submitted", "in_review", "revisions", "published"}
    if payload.status not in valid_statuses:
        raise HTTPException(status_code=422, detail=f"status must be one of {sorted(valid_statuses)}")

    db_project = get_project(db, project_id=project_id, user_id=current_user.id)
    if not db_project:
        raise HTTPException(status_code=404, detail="Project not found")

    db_project.status = payload.status
    if payload.target_journal is not None:
        db_project.target_journal = payload.target_journal
    db.commit()
    db.refresh(db_project)
    return db_project

@router.delete("/{project_id}")
def remove_project(
    project_id: int,
    db: Session = Depends(get_db),
    current_user = Depends(get_current_user)
):
    ok = delete_project(db, project_id=project_id, user_id=current_user.id)
    if not ok:
        raise HTTPException(status_code=404, detail="Project not found")
    return {"deleted": True}

@router.post("/export")
def export_document(request: ExportRequest, current_user = Depends(get_current_user)):
    title = request.title
    content = request.content  # real HTML from the editor, not stripped plain text
    fmt = request.format.lower()

    if "docx" in fmt:
        doc = docx.Document()
        doc.add_heading(title, 0)
        html_to_docx(doc, content)

        file_stream = io.BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)

        return Response(
            content=file_stream.read(),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f'attachment; filename="{title}.docx"'}
        )

    elif "pdf" in fmt:
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)
        try:
            # write_html understands headings, bold/italic, lists, and tables —
            # far more faithful to the actual document than dumping plain text.
            safe_title = title.replace("<", "&lt;").replace(">", "&gt;")
            pdf.write_html(f"<h1>{safe_title}</h1>{content}")
        except Exception:
            # Fail-soft: if any HTML in the document trips up the renderer,
            # fall back to plain readable text rather than a failed export.
            pdf2 = FPDF()
            pdf2.add_page()
            pdf2.set_font("Helvetica", size=12)
            pdf2.cell(0, 10, text=title, new_x="LMARGIN", new_y="NEXT", align="C")
            safe_content = html_to_plain_text(content).encode("latin-1", "replace").decode("latin-1")
            pdf2.multi_cell(0, 10, text=safe_content)
            pdf = pdf2

        output = pdf.output()
        pdf_bytes = output.encode("latin-1") if isinstance(output, str) else bytes(output)

        return Response(
            content=pdf_bytes,
            media_type="application/pdf",
            headers={"Content-Disposition": f'attachment; filename="{title}.pdf"'}
        )

    elif "markdown" in fmt:
        md_content = f"# {title}\n\n{html_to_markdown(content)}"
        return Response(
            content=md_content.encode('utf-8'),
            media_type="text/markdown",
            headers={"Content-Disposition": f'attachment; filename="{title}.md"'}
        )

    else:
        # BibTeX — built from the real citations used in this document (passed
        # from the workspace's Cite feature), not a fabricated placeholder entry.
        if request.citations:
            entries = []
            for c in request.citations:
                if c.bibtex:
                    entries.append(c.bibtex)
            bib_content = "\n\n".join(entries) if entries else _empty_bibtex_note(title)
        else:
            bib_content = _empty_bibtex_note(title)

        return Response(
            content=bib_content.encode("utf-8"),
            media_type="application/x-bibtex",
            headers={"Content-Disposition": f'attachment; filename="{_slugify(title)}.bib"'}
        )


def _empty_bibtex_note(title: str) -> str:
    return (
        f"% No citations were inserted in \"{title}\" via the workspace's Cite tool.\n"
        f"% Use the Cite button in the editor toolbar to add references, then export again."
    )


def _slugify(title: str) -> str:
    slug = re.sub(r"[^a-zA-Z0-9]+", "-", title or "document").strip("-").lower()
    return slug or "document"

@router.get("/")
def list_projects(db: Session = Depends(get_db), current_user = Depends(get_current_user)):
    projects = get_projects(db, user_id=current_user.id)
    result = []
    for p in projects:
        chapter = db.query(ChapterModel).filter(ChapterModel.project_id == p.id).first()
        word_count = len(chapter.content.split()) if chapter and chapter.content else 0
        result.append({
            "id": p.id,
            "title": p.title,
            "description": p.description,
            "field": p.field,
            "status": p.status,
            "target_journal": p.target_journal,
            "word_count": word_count,
            "updated_at": p.updated_at.isoformat() if p.updated_at else None,
        })
    return result

@router.get("/{project_id}")
def read_project(project_id: int, db: Session = Depends(get_db), current_user = Depends(get_current_user)):
    db_project = get_project(db, project_id=project_id, user_id=current_user.id)
    if db_project is None:
        raise HTTPException(status_code=404, detail="Project not found")
    
    # Grab the saved text to send back to the editor
    db_chapter = db.query(ChapterModel).filter(ChapterModel.project_id == project_id).first()
    content = db_chapter.content if db_chapter else ""
    
    return {
        "id": db_project.id,
        "title": db_project.title,
        "description": db_project.description,
        "content": content
    }

@router.post("/", response_model=ProjectResponse)
def add_project(project: ProjectCreate, db: Session = Depends(get_db), current_user = Depends(get_current_user)):
    return create_project(db, project=project, user_id=current_user.id)