"""
Converts the workspace editor's HTML content into a properly structured
DOCX (via python-docx) or Markdown string — replacing the previous approach,
which literally dumped raw HTML tags as visible text in exported documents.
"""
from html.parser import HTMLParser

import docx


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

class _DocxBuilder(HTMLParser):
    """Walks the editor's HTML and writes real, structured content into a
    python-docx Document — headings, paragraphs, bold/italic/underline runs,
    bullet/numbered lists, tables, blockquotes, and links (as visible text)."""

    def __init__(self, doc: docx.Document):
        super().__init__(convert_charrefs=True)
        self.doc = doc
        self.bold = False
        self.italic = False
        self.underline = False
        self.list_stack = []
        self.current_para = None
        self.in_table = False
        self.table_rows = []
        self.current_row = None
        self._active_cell = None
        self.skip_depth = 0

    def _ensure_paragraph(self):
        if self.current_para is None:
            self.current_para = self.doc.add_paragraph()
        return self.current_para

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self.skip_depth += 1
            return

        if tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            level = int(tag[1])
            self.current_para = self.doc.add_heading(level=min(level, 4))
        elif tag == "p" and not self.in_table:
            self.current_para = self.doc.add_paragraph()
        elif tag == "br":
            if not self.in_table:
                p = self._ensure_paragraph()
                p.add_run().add_break()
        elif tag in ("strong", "b"):
            self.bold = True
        elif tag in ("em", "i"):
            self.italic = True
        elif tag == "u":
            self.underline = True
        elif tag in ("ul", "ol"):
            self.list_stack.append(tag)
        elif tag == "li" and not self.in_table:
            style = "List Bullet" if (self.list_stack and self.list_stack[-1] == "ul") else "List Number"
            self.current_para = self.doc.add_paragraph(style=style)
        elif tag == "blockquote" and not self.in_table:
            self.current_para = self.doc.add_paragraph(style="Intense Quote")
        elif tag == "table":
            self.in_table = True
            self.table_rows = []
        elif tag == "tr":
            self.current_row = []
        elif tag in ("td", "th"):
            cell_data = {"is_header": tag == "th", "runs": []}
            self.current_row.append(cell_data)
            self._active_cell = cell_data

    def handle_endtag(self, tag):
        if tag in ("script", "style"):
            self.skip_depth = max(0, self.skip_depth - 1)
            return

        if tag in ("strong", "b"):
            self.bold = False
        elif tag in ("em", "i"):
            self.italic = False
        elif tag == "u":
            self.underline = False
        elif tag in ("ul", "ol"):
            if self.list_stack:
                self.list_stack.pop()
        elif tag in ("p", "li", "blockquote", "h1", "h2", "h3", "h4", "h5", "h6") and not self.in_table:
            self.current_para = None
        elif tag == "tr":
            if self.current_row is not None:
                self.table_rows.append(self.current_row)
            self.current_row = None
        elif tag in ("td", "th"):
            self._active_cell = None
        elif tag == "table":
            self._flush_table()
            self.in_table = False

    def _flush_table(self):
        if not self.table_rows:
            return
        n_cols = max(len(r) for r in self.table_rows)
        n_rows = len(self.table_rows)
        table = self.doc.add_table(rows=n_rows, cols=n_cols)
        try:
            table.style = "Table Grid"
        except KeyError:
            pass

        for r, row in enumerate(self.table_rows):
            for c in range(n_cols):
                cell = table.cell(r, c)
                if c < len(row):
                    cell_data = row[c]
                    p = cell.paragraphs[0]
                    for text, bold, italic, underline in cell_data["runs"]:
                        if not text:
                            continue
                        run = p.add_run(text)
                        run.bold = bold or cell_data["is_header"]
                        run.italic = italic
                        run.underline = underline
        self.doc.add_paragraph()

    def handle_data(self, data):
        if self.skip_depth or not data:
            return

        if self.in_table:
            if self._active_cell is not None:
                self._active_cell["runs"].append((data, self.bold, self.italic, self.underline))
            return

        p = self._ensure_paragraph()
        run = p.add_run(data)
        run.bold = self.bold
        run.italic = self.italic
        run.underline = self.underline


def html_to_docx(doc: docx.Document, html: str):
    """Populates a python-docx Document with the editor's HTML content,
    properly structured (not dumped as raw tag text)."""
    builder = _DocxBuilder(doc)
    builder.feed(html or "")
    builder.close()


# ---------------------------------------------------------------------------
# Markdown
# ---------------------------------------------------------------------------

class _MarkdownBuilder(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.out = []
        self.list_stack = []
        self.li_index = []
        self.skip_depth = 0
        self.in_table = False
        self.table_rows = []
        self.current_row = None
        self.current_cell_text = []

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self.skip_depth += 1
        elif tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            self.out.append("\n" + "#" * int(tag[1]) + " ")
        elif tag == "p":
            self.out.append("\n\n")
        elif tag == "br":
            self.out.append("  \n")
        elif tag in ("strong", "b"):
            self.out.append("**")
        elif tag in ("em", "i"):
            self.out.append("*")
        elif tag == "ul":
            self.list_stack.append("ul")
        elif tag == "ol":
            self.list_stack.append("ol")
            self.li_index.append(0)
        elif tag == "li":
            indent = "  " * max(0, len(self.list_stack) - 1)
            if self.list_stack and self.list_stack[-1] == "ol":
                self.li_index[-1] += 1
                self.out.append(f"\n{indent}{self.li_index[-1]}. ")
            else:
                self.out.append(f"\n{indent}- ")
        elif tag == "blockquote":
            self.out.append("\n> ")
        elif tag == "table":
            self.in_table = True
            self.table_rows = []
        elif tag == "tr":
            self.current_row = []
        elif tag in ("td", "th"):
            self.current_cell_text = []
        elif tag == "hr":
            self.out.append("\n\n---\n\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style"):
            self.skip_depth = max(0, self.skip_depth - 1)
        elif tag in ("strong", "b"):
            self.out.append("**")
        elif tag in ("em", "i"):
            self.out.append("*")
        elif tag == "ul":
            if self.list_stack:
                self.list_stack.pop()
        elif tag == "ol":
            if self.list_stack:
                self.list_stack.pop()
            if self.li_index:
                self.li_index.pop()
        elif tag in ("td", "th"):
            if self.current_row is not None:
                self.current_row.append("".join(self.current_cell_text).strip())
        elif tag == "tr":
            if self.current_row is not None:
                self.table_rows.append(self.current_row)
            self.current_row = None
        elif tag == "table":
            self._flush_table()
            self.in_table = False

    def _flush_table(self):
        if not self.table_rows:
            return
        self.out.append("\n\n")
        header, *rest = self.table_rows
        self.out.append("| " + " | ".join(header) + " |\n")
        self.out.append("| " + " | ".join(["---"] * len(header)) + " |\n")
        for row in rest:
            self.out.append("| " + " | ".join(row) + " |\n")

    def handle_data(self, data):
        if self.skip_depth:
            return
        if self.in_table and self.current_row is not None:
            self.current_cell_text.append(data)
            return
        self.out.append(data)

    def result(self) -> str:
        text = "".join(self.out)
        while "\n\n\n" in text:
            text = text.replace("\n\n\n", "\n\n")
        return text.strip()


def html_to_markdown(html: str) -> str:
    builder = _MarkdownBuilder()
    builder.feed(html or "")
    builder.close()
    return builder.result()


# ---------------------------------------------------------------------------
# Plain text (simple fallback / used as PDF safety net)
# ---------------------------------------------------------------------------

class _PlainTextBuilder(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.out = []
        self.skip_depth = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self.skip_depth += 1
        elif tag in ("p", "br", "div", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6"):
            self.out.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style"):
            self.skip_depth = max(0, self.skip_depth - 1)

    def handle_data(self, data):
        if not self.skip_depth:
            self.out.append(data)

    def result(self) -> str:
        text = "".join(self.out)
        while "\n\n\n" in text:
            text = text.replace("\n\n\n", "\n\n")
        return text.strip()


def html_to_plain_text(html: str) -> str:
    builder = _PlainTextBuilder()
    builder.feed(html or "")
    builder.close()
    return builder.result()
